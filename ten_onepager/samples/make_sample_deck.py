"""Generate the FICTIONAL sample deck used for demos and tests.

Northpeak Cold Chain does not exist; every name and figure is invented. The
deck deliberately includes things the review screen should catch: two
different ARR figures, pilots alongside paying customers, soft-circled money
alongside cash received, a projection, and no exit information.

    python samples/make_sample_deck.py            # writes sample_deck.pptx
    python samples/make_sample_deck.py --pdf --ppt  # also converts with LibreOffice
"""

from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

FICTION = "FICTIONAL SAMPLE: Northpeak Cold Chain is an invented company. All names and figures are made up."

SLIDES = [
    ("Northpeak Cold Chain",
     ["Real-time spoilage prevention for regional food distributors", FICTION]),
    ("The problem",
     ["US regional food distributors lose 4-6% of perishable inventory to temperature excursions "
      "(Northpeak survey of 40 distributors, 2025).",
      "A typical customer: $180M annual revenue and $3.6M in annual spoilage write-offs.",
      "Existing data loggers report excursions after the fact, when product is already lost."]),
    ("Our solution",
     ["Wireless sensors plus prediction software that alerts staff about 6 hours before an excursion.",
      "Pilot customers cut spoilage by 31% over 6 months.",
      "Installs in one day per warehouse; no changes to existing refrigeration."]),
    ("Traction",
     ["ARR: $1.2M (June 2026)",
      "14 paying customers; 5 additional pilots in progress",
      "Net revenue retention: 118% (trailing 12 months)",
      "3 signed letters of intent from national grocery chains"]),
    ("Market",
     ["TAM: $9.4B global cold-chain monitoring market (2026, third-party industry report)",
      "SAM: $1.1B US regional food distributors (bottom-up: 2,600 distributors x $420K annual spoilage-prevention budget)"]),
    ("Why we win",
     ["Predictive alerts versus after-the-fact logging",
      "2 US patents pending on the excursion-prediction model",
      "Trained on 3.2 billion sensor readings from customer sites"]),
    ("Team",
     ["Dana Whitfield, CEO: former VP Operations at a regional food distributor (12 years)",
      "Ravi Menon, CTO: IoT engineering lead; 2 prior startups, 1 acquired",
      "Head of Sales: hiring"]),
    ("The raise",
     ["Raising $4.0M Seed on a SAFE with a $16M post-money valuation cap",
      "$0.6M received from angel investors",
      "$1.5M additionally soft-circled"]),
    ("Use of funds",
     ["Product & engineering 40% | Sales & marketing 35% | Operations 15% | G&A 10%",
      "18-month runway",
      "Milestone: $3.5M ARR by end of 2027 (projected)"]),
    ("Financial outlook",
     ["Current ARR: $1.4M",
      "Projected revenue: 2027E $3.5M, 2028E $8.0M",
      "Gross margin target: 70%"]),
]


def build(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[1]
    for i, (title, bullets) in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1]
        body.left, body.top, body.width, body.height = Inches(0.7), Inches(1.6), Inches(7.4 if i == 4 else 12), Inches(5)
        tf = body.text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(20)
        if i == 4:
            data = CategoryChartData()
            data.categories = ["Q3 2025", "Q4 2025", "Q1 2026", "Q2 2026"]
            data.add_series("ARR ($M)", (0.45, 0.68, 0.95, 1.2))
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(8.3), Inches(1.6),
                                   Inches(4.5), Inches(4.2), data)
            slide.notes_slide.notes_text_frame.text = "Pilots convert to paid after 90 days on average."
        if i == 1:
            slide.notes_slide.notes_text_frame.text = FICTION
    prs.save(path)
    return path


def convert(src: Path, fmt: str) -> Path:
    from onepager.extraction import find_soffice

    soffice = find_soffice()
    if not soffice:
        raise SystemExit("LibreOffice not found; cannot convert.")
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp, "profile").resolve().as_uri()
        subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--convert-to", fmt,
                        "--outdir", str(src.parent), str(src)], check=True, capture_output=True, timeout=180)
    return src.with_suffix("." + fmt)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--pdf", action="store_true")
    ap.add_argument("--ppt", action="store_true")
    args = ap.parse_args()
    out = build(HERE / "sample_deck.pptx")
    print("wrote", out)
    if args.pdf:
        print("wrote", convert(out, "pdf"))
    if args.ppt:
        print("wrote", convert(out, "ppt"))
