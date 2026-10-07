"""Deck text extraction with page/slide references.

Supported inputs:
  .pdf   PyMuPDF text extraction; optional OCR for scanned (image-only) pages
         through PyMuPDF's Tesseract integration.
  .pptx  python-pptx: text frames, grouped shapes, tables, chart data and
         speaker notes.
  .ppt   converted to .pptx with LibreOffice when it is installed.

Every page keeps a reference ("Page 3" / "Slide 3") so the analysis can cite
where each figure came from.
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".pptx", ".ppt")

# A page with fewer characters than this, but with images, is treated as scanned.
LOW_TEXT_CHARS = 25

PPT_CONVERSION_HELP = (
    "Legacy .ppt files need LibreOffice for conversion, and it was not found. "
    "Either install LibreOffice (https://www.libreoffice.org) and retry, or open the "
    "file in PowerPoint and use File > Save As > PowerPoint Presentation (.pptx), "
    "or File > Export > PDF, then upload the converted file."
)


class ExtractionError(Exception):
    """The file could not be read; the message is safe to show to the user."""


@dataclass
class DeckPage:
    ref: str            # "Slide 3" or "Page 3"
    number: int
    text: str
    ocr: bool = False
    low_text: bool = False


@dataclass
class DeckContent:
    filename: str
    source_type: str    # "pdf", "pptx", "ppt"
    pages: list[DeckPage]
    warnings: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)

    def to_prompt_text(self) -> str:
        return "\n\n".join(f"=== {p.ref} ===\n{p.text or '[no extractable text]'}" for p in self.pages)


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #
def extract_deck(path: str | os.PathLike, *, ocr: bool = False, display_name: str | None = None) -> DeckContent:
    path = Path(path)
    name = display_name or path.name
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(f"Unsupported file type '{ext}'. Upload a PDF, PPTX or PPT file.")
    if not path.is_file():
        raise ExtractionError(f"File not found: {name}")

    if ext == ".pdf":
        deck = _extract_pdf(path, name, ocr=ocr)
    elif ext == ".pptx":
        deck = _extract_pptx(path, name)
    else:
        with tempfile.TemporaryDirectory(prefix="onepager_ppt_") as tmp:
            converted = convert_ppt_to_pptx(path, Path(tmp))
            deck = _extract_pptx(converted, name)
        deck.source_type = "ppt"
        deck.warnings.insert(0, "Converted from legacy .ppt with LibreOffice; check that charts and tables came through.")

    if not any(p.text.strip() for p in deck.pages):
        deck.warnings.append("No readable text was found in this file.")
    return deck


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def ocr_status() -> tuple[bool, str]:
    """Whether OCR can run, with a user-facing explanation."""
    import pymupdf

    # PyMuPDF ships the Tesseract engine; it only needs the language data.
    try:
        tessdata = pymupdf.get_tessdata()
    except Exception:  # noqa: BLE001 - raised when Tesseract language data is absent
        tessdata = None
    if tessdata and os.path.isdir(tessdata):
        return True, "OCR available (Tesseract)."
    return False, (
        "OCR is unavailable because Tesseract is not installed. Install Tesseract OCR "
        "(https://github.com/tesseract-ocr/tesseract) and make sure it is on PATH or "
        "TESSDATA_PREFIX points to its tessdata folder."
    )


def _extract_pdf(path: Path, name: str, *, ocr: bool) -> DeckContent:
    import pymupdf

    try:
        doc = pymupdf.open(path)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Could not open the PDF: {exc}") from exc
    if doc.needs_pass:
        raise ExtractionError("This PDF is password-protected. Remove the password and upload it again.")

    can_ocr, ocr_msg = ocr_status() if ocr else (False, "")
    pages, scanned = [], []
    with doc:
        for i, page in enumerate(doc, start=1):
            text = _clean_pdf_text(page.get_text("text"))
            low = len(text) < LOW_TEXT_CHARS and bool(page.get_images())
            used_ocr = False
            if low and ocr and can_ocr:
                try:
                    tp = page.get_textpage_ocr(full=True, dpi=200)
                    text = _clean_pdf_text(page.get_text("text", textpage=tp))
                    used_ocr = True
                except Exception:  # noqa: BLE001 - keep going with whatever text exists
                    pass
            if low and not used_ocr:
                scanned.append(i)
            pages.append(DeckPage(f"Page {i}", i, text, ocr=used_ocr, low_text=low and not used_ocr))

    deck = DeckContent(name, "pdf", pages)
    if scanned:
        listed = ", ".join(map(str, scanned[:12])) + ("..." if len(scanned) > 12 else "")
        if not ocr:
            deck.warnings.append(f"Pages {listed} look scanned (images with little text). Enable OCR to read them.")
        elif not can_ocr:
            deck.warnings.append(f"Pages {listed} look scanned. {ocr_msg}")
        else:
            deck.warnings.append(f"OCR could not read pages {listed}.")
    if any(p.ocr for p in pages):
        deck.warnings.append("Some pages were read with OCR; check figures on those pages carefully.")
    return deck


_PRIVATE_USE = re.compile("[-�]")
_BULLET_ONLY = re.compile(r"^[\s•●▪■◦\-–*]*$")


def _clean_pdf_text(text: str) -> str:
    """Drop private-use glyphs (icon-font bullets) and bullet-only lines."""
    lines = [_PRIVATE_USE.sub("", ln).rstrip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if not _BULLET_ONLY.match(ln)).strip()


# --------------------------------------------------------------------------- #
# PPTX
# --------------------------------------------------------------------------- #
def _extract_pptx(path: Path, name: str) -> DeckContent:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Could not open the PowerPoint file: {exc}") from exc

    pages, image_only = [], []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        has_picture = False
        for shape in slide.shapes:
            has_picture |= _collect_shape(shape, parts)
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                parts.append(f"[Speaker notes] {notes}")
        text = "\n".join(parts).strip()
        low = len(text) < LOW_TEXT_CHARS and has_picture
        if low:
            image_only.append(i)
        pages.append(DeckPage(f"Slide {i}", i, text, low_text=low))

    deck = DeckContent(name, "pptx", pages)
    if image_only:
        deck.warnings.append(
            f"Slides {', '.join(map(str, image_only))} contain images with little text. Text inside images "
            "is not read from PowerPoint files; export the deck to PDF and enable OCR to include it."
        )
    return deck


def _collect_shape(shape, parts: list[str]) -> bool:
    """Append a shape's text to parts. Returns True if the shape is a picture."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        found = False
        for sub in shape.shapes:
            found |= _collect_shape(sub, parts)
        return found
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip()).strip()
        if text:
            parts.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows]
        rows = [r for r in rows if r.replace("|", "").strip()]
        if rows:
            parts.append("[Table]\n" + "\n".join(rows))
    if getattr(shape, "has_chart", False) and shape.has_chart:
        parts.append(_chart_text(shape.chart))
    return False


def _chart_text(chart) -> str:
    lines = ["[Chart]"]
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            lines[0] = f"[Chart: {chart.chart_title.text_frame.text.strip()}]"
    except Exception:  # noqa: BLE001
        pass
    try:
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        if cats:
            lines.append("Categories: " + ", ".join(cats))
        for series in plot.series:
            vals = ", ".join("" if v is None else f"{v:g}" for v in series.values)
            lines.append(f"{series.name}: {vals}")
    except Exception:  # noqa: BLE001 - charts without readable data
        lines.append("(chart data not readable)")
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# Legacy PPT
# --------------------------------------------------------------------------- #
def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if os.path.isfile(candidate):
            return candidate
    return None


def convert_ppt_to_pptx(path: Path, out_dir: Path) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise ExtractionError(PPT_CONVERSION_HELP)
    # A private profile directory avoids clashing with a running LibreOffice.
    profile = (out_dir / "lo_profile").resolve().as_uri()
    cmd = [soffice, f"-env:UserInstallation={profile}", "--headless",
           "--convert-to", "pptx", "--outdir", str(out_dir), str(path)]
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=180)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as exc:
        raise ExtractionError(f"LibreOffice could not convert the .ppt file ({type(exc).__name__}). "
                              + PPT_CONVERSION_HELP) from exc
    converted = out_dir / (path.stem + ".pptx")
    if not converted.is_file():
        raise ExtractionError("LibreOffice did not produce a .pptx file. " + PPT_CONVERSION_HELP)
    return converted
