from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

from onepager.extraction import ExtractionError, extract_deck, find_soffice


def _pptx(path: Path) -> Path:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Traction"
    s1.placeholders[1].text_frame.text = "ARR $2.5M as of March 2026"
    s1.notes_slide.notes_text_frame.text = "Churn under 3%"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Funds"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Category", "Share"
    table.cell(1, 0).text, table.cell(1, 1).text = "R&D", "45%"
    data = CategoryChartData()
    data.categories = ["2024", "2025"]
    data.add_series("Revenue ($M)", (0.8, 1.9))
    s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(2), Inches(4), Inches(3), data)
    prs.save(path)
    return path


def _pdf(path: Path, pages: list[str], image_page: bool = False) -> Path:
    c = Canvas(str(path), pagesize=letter)
    for text in pages:
        c.drawString(72, 700, text)
        c.showPage()
    if image_page:
        from PIL import Image

        img = Image.new("RGB", (400, 200), "white")
        c.drawImage(ImageReader(img), 72, 400, 400, 200)
        c.showPage()
    c.save()
    return path


def test_pptx_text_tables_charts_notes_with_slide_refs(tmp_path):
    deck = extract_deck(_pptx(tmp_path / "d.pptx"))
    assert deck.source_type == "pptx"
    assert [p.ref for p in deck.pages] == ["Slide 1", "Slide 2"]
    assert "ARR $2.5M" in deck.pages[0].text
    assert "[Speaker notes] Churn under 3%" in deck.pages[0].text
    assert "R&D | 45%" in deck.pages[1].text
    assert "Revenue ($M): 0.8, 1.9" in deck.pages[1].text
    assert "=== Slide 2 ===" in deck.to_prompt_text()


def test_pdf_text_with_page_refs(tmp_path):
    deck = extract_deck(_pdf(tmp_path / "d.pdf", ["Raising $3M seed", "TAM $12B global"]))
    assert deck.source_type == "pdf"
    assert [p.ref for p in deck.pages] == ["Page 1", "Page 2"]
    assert "TAM $12B" in deck.pages[1].text
    assert not deck.warnings


def test_scanned_page_is_flagged_without_ocr(tmp_path):
    deck = extract_deck(_pdf(tmp_path / "d.pdf", ["Intro page with enough text to count"], image_page=True))
    assert deck.pages[1].low_text
    assert any("Enable OCR" in w for w in deck.warnings)


def test_ocr_requested_but_unavailable_is_explained(tmp_path, monkeypatch):
    import onepager.extraction as ex

    monkeypatch.setattr(ex, "ocr_status", lambda: (False, "OCR is unavailable because Tesseract is not installed."))
    deck = extract_deck(_pdf(tmp_path / "d.pdf", ["Some text on page one here"], image_page=True), ocr=True)
    assert any("Tesseract is not installed" in w for w in deck.warnings)


def test_unsupported_type_rejected(tmp_path):
    f = tmp_path / "deck.key"
    f.write_text("x")
    with pytest.raises(ExtractionError, match="Unsupported"):
        extract_deck(f)


def test_ppt_without_libreoffice_explains_conversion(tmp_path, monkeypatch):
    import onepager.extraction as ex

    monkeypatch.setattr(ex, "find_soffice", lambda: None)
    f = tmp_path / "old.ppt"
    f.write_bytes(b"not really a ppt")
    with pytest.raises(ExtractionError, match="Save As"):
        extract_deck(f)


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_ppt_converted_with_libreoffice(samples_dir):
    deck = extract_deck(samples_dir / "sample_deck.ppt")
    assert deck.source_type == "ppt"
    assert len(deck.pages) == 10
    assert "ARR: $1.2M" in deck.pages[3].text


def test_sample_pdf_and_pptx_agree(samples_dir):
    pdf = extract_deck(samples_dir / "sample_deck.pdf")
    pptx = extract_deck(samples_dir / "sample_deck.pptx")
    assert len(pdf.pages) == len(pptx.pages) == 10
    for needle in ("$4.0M", "$9.4B", "14 paying customers"):
        assert any(needle in p.text for p in pdf.pages)
        assert any(needle in p.text for p in pptx.pages)
    assert "�" not in pdf.to_prompt_text()
