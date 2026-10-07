"""Stage 1 — Extraction.

Detect the input file type by extension and pull out all text and tables,
preserving slide / page / table structure with human-readable markers so the
analysis stage can reason about where each figure came from.
"""

from __future__ import annotations

import os


class UnsupportedFileError(ValueError):
    """Raised when the input file extension is not one we know how to read."""


def extract(path: str) -> str:
    """Route to the correct extractor based on file extension.

    Returns a single concatenated text blob with structure markers.
    Raises UnsupportedFileError for anything we can't parse.
    """
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Input file not found: {path}")

    ext = os.path.splitext(path)[1].lower()
    if ext == ".pptx":
        return _extract_pptx(path)
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    raise UnsupportedFileError(
        f"Unsupported file type '{ext}'. Supported types: .pptx, .pdf, .docx"
    )


def _table_to_text(rows: list[list[str]]) -> str:
    """Render a table (list of rows of cells) as pipe-delimited text lines."""
    lines = []
    for row in rows:
        cells = [(c if c is not None else "").strip() for c in row]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(
                    p.text for p in shape.text_frame.paragraphs if p.text
                ).strip()
                if text:
                    parts.append(text)
            if shape.has_table:
                rows = [
                    [cell.text for cell in row.cells]
                    for row in shape.table.rows
                ]
                table_text = _table_to_text(rows)
                if table_text:
                    parts.append("[Table]\n" + table_text)
        # Speaker notes often carry useful context.
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                parts.append("[Notes] " + notes)
    return "\n".join(parts).strip()


def _extract_pdf(path: str) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            parts.append(f"--- Page {i} ---")
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(text)
            for table in page.extract_tables():
                table_text = _table_to_text(table)
                if table_text:
                    parts.append("[Table]\n" + table_text)
    return "\n".join(parts).strip()


def _extract_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts: list[str] = ["--- Document ---"]
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for t, table in enumerate(doc.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        table_text = _table_to_text(rows)
        if table_text:
            parts.append(f"[Table {t}]\n" + table_text)
    return "\n".join(parts).strip()
